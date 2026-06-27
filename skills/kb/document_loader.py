"""知识库文档解析。

- PDF / DOCX / TXT / MD：轻量库（pypdf、python-docx），无 NLTK 依赖
- 旧版 .doc：unstructured + NLTK 分词数据（仅 .doc 需要，建议安装时预置）
"""
from __future__ import annotations

import os
import ssl
import zipfile
from pathlib import Path

from loguru import logger

from utils import CONFIG

_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gbk", "gb18030", "latin-1")
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# unstructured 解析 .doc 时可能调用的 NLTK 数据包（pip 不含数据，需单独下载或离线目录）
_NLTK_RESOURCES = (
    "punkt_tab",
    "punkt",
    "averaged_perceptron_tagger_eng",
    "averaged_perceptron_tagger",
)
_nltk_ready = False
_nltk_paths_inited = False


def _read_plain_text(path: str) -> str:
    last_err: Exception | None = None
    for enc in _TEXT_ENCODINGS:
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError as exc:
            last_err = exc
    raise ValueError(f"无法解码文本文件: {path}") from last_err


def _load_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n".join(parts)


def _load_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _nltk_resource_paths(name: str) -> tuple[str, ...]:
    if name.startswith("punkt"):
        return (f"tokenizers/{name}",)
    if "tagger" in name:
        return (f"taggers/{name}", f"taggers/{name}/{name}")
    return (name,)


def _nltk_resource_exists(name: str) -> bool:
    import nltk

    for path in _nltk_resource_paths(name):
        try:
            nltk.data.find(path)
            return True
        except LookupError:
            continue
    return False


def _init_nltk_paths() -> None:
    """注册 NLTK 数据目录（环境变量 NLTK_DATA 或 config kb.nltk_data_dir）。"""
    global _nltk_paths_inited
    if _nltk_paths_inited:
        return
    import nltk

    candidates: list[str] = []
    env_dir = os.environ.get("NLTK_DATA", "").strip()
    if env_dir:
        candidates.append(env_dir)
    cfg_dir = str((CONFIG.get("kb") or {}).get("nltk_data_dir") or "").strip()
    if cfg_dir:
        candidates.append(cfg_dir)
    # 项目内离线目录（安装脚本可预置）
    bundled = Path(__file__).resolve().parents[2] / "data" / "nltk_data"
    if bundled.is_dir():
        candidates.append(str(bundled))

    for d in candidates:
        if d and os.path.isdir(d) and d not in nltk.data.path:
            nltk.data.path.insert(0, d)
            logger.info(f"NLTK 数据目录: {d}")
    _nltk_paths_inited = True


def _download_nltk_resource(name: str) -> bool:
    import nltk

    # 内网/自签证书：NLTK 官方源常 SSL 失败，下载时跳过证书校验
    prev = ssl._create_default_https_context
    try:
        ssl._create_default_https_context = ssl._create_unverified_context
        return bool(nltk.download(name, quiet=True))
    except Exception as exc:
        logger.warning(f"nltk.download({name}) 失败: {exc}")
        return False
    finally:
        ssl._create_default_https_context = prev


def _ensure_nltk_data() -> None:
    """确保 unstructured 解析 .doc 所需的 NLTK 数据已就绪。"""
    global _nltk_ready
    if _nltk_ready:
        return

    try:
        import nltk  # noqa: F401
    except ImportError as exc:
        raise RuntimeError(
            "解析 .doc 需要 nltk，请执行: pip install nltk"
        ) from exc

    _init_nltk_paths()

    missing = [n for n in _NLTK_RESOURCES if not _nltk_resource_exists(n)]
    if not missing:
        _nltk_ready = True
        return

    logger.warning(
        f"本地未找到 NLTK 数据 {missing}，尝试联网下载（仅 .doc 需要；"
        f"建议安装时执行: python -m skills.kb.setup_nltk）"
    )
    for name in missing:
        logger.info(f"下载 NLTK 资源: {name}")
        if not _download_nltk_resource(name):
            raise RuntimeError(
                f"NLTK 资源 {name} 不可用。请在可联网机器下载后拷贝到 NLTK_DATA，"
                f"或执行: python -m skills.kb.setup_nltk"
            )

    _nltk_ready = True


def warmup_nltk_data() -> None:
    """启动时后台预检 NLTK 数据（不阻塞服务；.doc 上传前尽量已就绪）。"""
    try:
        _ensure_nltk_data()
        logger.info("NLTK 数据已就绪（.doc 解析可用）")
    except Exception as exc:
        logger.warning(f"NLTK 数据未就绪，.doc 上传可能失败: {exc}")


def _load_unstructured(path: str) -> str:
    """旧版 .doc 等格式：unstructured 解析（依赖 NLTK）。"""
    _ensure_nltk_data()
    try:
        from unstructured.partition.auto import partition
    except ImportError as exc:
        raise RuntimeError(
            "解析 .doc 需要 unstructured，请执行: pip install 'unstructured[doc,docx,pdf]'"
        ) from exc

    elements = partition(filename=path)
    parts: list[str] = []
    for el in elements:
        t = getattr(el, "text", None) or str(el)
        t = (t or "").strip()
        if t:
            parts.append(t)
    return "\n".join(parts)


def _file_magic(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read(8)


def _load_by_magic(path: str) -> str | None:
    """按文件头识别 PDF/DOCX/旧版 DOC（扩展名错误时兜底）。"""
    head = _file_magic(path)
    if head.startswith(b"%PDF"):
        return _load_pdf(path)
    if head.startswith(b"PK"):
        try:
            return _load_docx(path)
        except Exception:
            if zipfile.is_zipfile(path):
                raise ValueError(
                    "检测到 ZIP/Office 文件，但无法作为 DOCX 解析，请确认格式"
                ) from None
    if head.startswith(_OLE_MAGIC):
        return _load_unstructured(path)
    return None


def load_document_text(file_path: str) -> str:
    """解析文档为纯文本。解析结果为空时抛出 ValueError。"""
    path = os.path.abspath(file_path)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"文件不存在: {path}")

    ext = Path(path).suffix.lower()
    logger.info(f"解析文档: {path} (type={ext or 'unknown'})")

    if ext in {".txt", ".text", ".md", ".markdown"}:
        text = _read_plain_text(path)
    elif ext == ".pdf":
        text = _load_pdf(path)
    elif ext == ".docx":
        text = _load_docx(path)
    elif ext == ".doc":
        text = _load_unstructured(path)
    else:
        magic_text = _load_by_magic(path)
        if magic_text is not None:
            text = magic_text
        else:
            try:
                text = _read_plain_text(path)
            except ValueError:
                text = _load_unstructured(path)

    text = text.strip()
    if not text:
        raise ValueError(f"文档解析为空: {file_path}")
    return text
